"""
DOCX resume builder for JobAgent V2.

Generates a clean, editable, single-column Microsoft Word (.docx) resume designed for ATS parsing.
"""

from pathlib import Path
from typing import Dict, Any, Union
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def build_docx_resume(
    resume_data: Dict[str, Any],
    output_path: Union[str, Path],
) -> Path:
    """
    Builds a single-column ATS-compliant .docx document from structured resume data.
    """
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    
    doc = Document()
    
    # Page setup: Standard letter, 0.6 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)
        
    # Styles helper
    def add_section_header(title: str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(title.upper())
        run.bold = True
        run.font.size = Pt(11)
        run.font.name = "Calibri"
        run.font.color.rgb = RGBColor(30, 41, 59) # Slate 800
        
    # 1. Header / Contact
    contact = resume_data.get("contact", {})
    name_p = doc.add_paragraph()
    name_p.paragraph_format.space_before = Pt(0)
    name_p.paragraph_format.space_after = Pt(2)
    name_run = name_p.add_run(contact.get("name", "CANDIDATE"))
    name_run.bold = True
    name_run.font.size = Pt(18)
    name_run.font.name = "Calibri"
    name_run.font.color.rgb = RGBColor(15, 23, 42) # Slate 900
    
    contact_parts = []
    if contact.get("email"):
        contact_parts.append(contact["email"])
    if contact.get("phone"):
        contact_parts.append(contact["phone"])
    if contact.get("location"):
        contact_parts.append(contact["location"])
    if contact.get("linkedin"):
        contact_parts.append(contact["linkedin"])
    if contact.get("github"):
        contact_parts.append(contact["github"])
        
    contact_p = doc.add_paragraph()
    contact_p.paragraph_format.space_before = Pt(0)
    contact_p.paragraph_format.space_after = Pt(6)
    c_run = contact_p.add_run(" | ".join(contact_parts))
    c_run.font.size = Pt(9.5)
    c_run.font.name = "Calibri"
    c_run.font.color.rgb = RGBColor(71, 85, 105)
    
    # 2. Summary
    summary_text = resume_data.get("summary", "")
    if summary_text:
        add_section_header("Professional Summary")
        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(0)
        sp.paragraph_format.space_after = Pt(4)
        s_run = sp.add_run(summary_text)
        s_run.font.size = Pt(10)
        s_run.font.name = "Calibri"

    # 3. Technical Skills
    skills = resume_data.get("skills", [])
    if skills:
        add_section_header("Technical Skills")
        sk_p = doc.add_paragraph()
        sk_p.paragraph_format.space_before = Pt(0)
        sk_p.paragraph_format.space_after = Pt(4)
        skill_names = [s.get("name", "") if isinstance(s, dict) else str(s) for s in skills]
        sk_run = sk_p.add_run(", ".join(skill_names))
        sk_run.font.size = Pt(10)
        sk_run.font.name = "Calibri"

    # 4. Experience
    exps = resume_data.get("experience", [])
    if exps:
        add_section_header("Professional Experience")
        for exp in exps:
            ep = doc.add_paragraph()
            ep.paragraph_format.space_before = Pt(4)
            ep.paragraph_format.space_after = Pt(1)
            
            role_run = ep.add_run(exp.get("role", "") + " — ")
            role_run.bold = True
            role_run.font.size = Pt(10.5)
            role_run.font.name = "Calibri"
            
            comp_run = ep.add_run(exp.get("company", ""))
            comp_run.font.size = Pt(10)
            comp_run.font.name = "Calibri"
            
            date_str = f" ({exp.get('start', '')} - {exp.get('end', '')})"
            date_run = ep.add_run(date_str)
            date_run.italic = True
            date_run.font.size = Pt(9.5)
            date_run.font.name = "Calibri"
            
            for ach in exp.get("achievements", []):
                ach_text = ach.get("text", "") if isinstance(ach, dict) else str(ach)
                bp = doc.add_paragraph(style='List Bullet')
                bp.paragraph_format.space_before = Pt(0)
                bp.paragraph_format.space_after = Pt(1.5)
                b_run = bp.add_run(ach_text)
                b_run.font.size = Pt(9.5)
                b_run.font.name = "Calibri"

    # 5. Projects
    projects = resume_data.get("projects", [])
    if projects:
        add_section_header("Projects")
        for p in projects:
            pp = doc.add_paragraph()
            pp.paragraph_format.space_before = Pt(3)
            pp.paragraph_format.space_after = Pt(1)
            
            pname_run = pp.add_run(p.get("name", ""))
            pname_run.bold = True
            pname_run.font.size = Pt(10)
            pname_run.font.name = "Calibri"
            
            techs = p.get("technologies", [])
            if techs:
                tech_run = pp.add_run(f" | Technologies: {', '.join(techs)}")
                tech_run.italic = True
                tech_run.font.size = Pt(9.5)
                tech_run.font.name = "Calibri"

    # 6. Education
    education = resume_data.get("education", [])
    if education:
        add_section_header("Education")
        for edu in education:
            edp = doc.add_paragraph()
            edp.paragraph_format.space_before = Pt(2)
            edp.paragraph_format.space_after = Pt(1)
            
            deg_run = edp.add_run(edu.get("degree", "") + " — ")
            deg_run.bold = True
            deg_run.font.size = Pt(10)
            deg_run.font.name = "Calibri"
            
            inst_run = edp.add_run(edu.get("institution", ""))
            inst_run.font.size = Pt(9.5)
            inst_run.font.name = "Calibri"
            
            dates = f" ({edu.get('start', '')} - {edu.get('end', '')})"
            dates_run = edp.add_run(dates)
            dates_run.italic = True
            dates_run.font.size = Pt(9)
            dates_run.font.name = "Calibri"

    doc.save(str(out))
    return out
